"""Pantallas de oficina: órdenes, expediente, orden impresa, fórmulas maestras y reporte."""
import io
from datetime import datetime

from flask import abort, current_app, flash, request, send_file, session, url_for

from nucleo.auth import login_required, requiere_facultad, requiere_modulo, usuario_puede
from nucleo.db import get_db_connection
from nucleo.vistas import full_page, hx_redirect, render_page

from .blueprint import ROLES_SUPERVISOR, accion, bp, firmar, usuario
from .servicios import ErrorOrdenes, booleano
from .servicios import expediente as srv_expediente
from .servicios import formulas as srv_formulas
from .servicios import ordenes as srv_ordenes

POR_PAGINA = 50
PESTANAS = {
    'proceso':    ('En proceso', ('EMITIDA', 'EN_PROCESO')),
    'revisar':    ('Por revisar', ('TERMINADA',)),
    'borradores': ('Borradores', ('BORRADOR',)),
    'cerradas':   ('Cerradas', ('REVISADA',)),
    'canceladas': ('Canceladas', ('CANCELADA',)),
}


def _pagina():
    try:
        return max(1, int(request.args.get('pagina', 1)))
    except ValueError:
        return 1


# ---------------------------------------------------------------------------
# Órdenes
# ---------------------------------------------------------------------------

@bp.route('')
@login_required
@requiere_modulo('ordenes')
def lista():
    from .ganchos import resumen
    pestana = request.args.get('pestana', 'proceso')
    if pestana not in PESTANAS:
        pestana = 'proceso'
    estatus = PESTANAS[pestana][1]
    pagina = _pagina()
    conn = get_db_connection()
    marcas = ','.join('?' * len(estatus))
    total = conn.execute(f'SELECT COUNT(*) FROM op_ordenes WHERE estatus IN ({marcas})', estatus).fetchone()[0]
    filas = conn.execute(
        'SELECT o.*, m.codigo AS producto_codigo, m.nombre AS producto_nombre, m.unidad, l.estatus AS granel_estatus '
        'FROM op_ordenes o JOIN materiales m ON m.id = o.material_id LEFT JOIN lotes l ON l.id = o.granel_lote_id '
        f'WHERE o.estatus IN ({marcas}) ORDER BY COALESCE(o.fecha_programada, o.creado_en) DESC, o.id DESC '
        'LIMIT ? OFFSET ?', (*estatus, POR_PAGINA, (pagina - 1) * POR_PAGINA)).fetchall()
    ordenes = []
    for f in filas:
        d = dict(f)
        d['avance'] = srv_ordenes.avance(conn, f['id']) if f['estatus'] in ('EMITIDA', 'EN_PROCESO') else (
            100 if f['estatus'] in ('TERMINADA', 'REVISADA') else 0)
        ordenes.append(d)
    conteos = {k: conn.execute(f'SELECT COUNT(*) FROM op_ordenes WHERE estatus IN ({",".join("?" * len(v[1]))})',
                               v[1]).fetchone()[0] for k, v in PESTANAS.items()}
    kpis = resumen(conn)
    conn.close()
    return render_page('ordenes/lista.html', ordenes=ordenes, pestana=pestana, pestanas=PESTANAS, conteos=conteos,
                       kpis=kpis, pagina=pagina, paginas=max(1, -(-total // POR_PAGINA)),
                       nombres=srv_ordenes.NOMBRE_ESTATUS)


@bp.route('/nueva', methods=['GET', 'POST'])
@login_required
@requiere_modulo('ordenes_planear')
def nueva():
    if request.method == 'POST':
        formula_id = request.form.get('formula_id', type=int)
        piezas = {int(k[7:]): v for k, v in request.form.items() if k.startswith('piezas_') and k[7:].isdigit()}

        def crear(conn):
            orden_id, avisos = srv_ordenes.crear(conn, formula_id=formula_id, tamano_lote=request.form.get('tamano_lote'),
                                                 fecha_programada=request.form.get('fecha_programada'),
                                                 piezas=piezas, observaciones=request.form.get('observaciones'),
                                                 usuario=usuario())
            return {'orden_id': orden_id, 'avisos': avisos}
        return accion(crear, 'Orden creada en borrador. Revísala y emítela.',
                      lambda r: url_for('ordenes.detalle', orden_id=r['orden_id']),
                      error_destino=url_for('ordenes.nueva', formula_id=formula_id or ''))
    conn = get_db_connection()
    formulas = conn.execute("SELECT f.id, f.version, f.tamano_lote, m.codigo, m.nombre, m.unidad FROM op_formulas f "
                            "JOIN materiales m ON m.id = f.material_id WHERE f.estatus = 'VIGENTE' ORDER BY m.codigo").fetchall()
    formula = disponibilidad = presentaciones = None
    tamano = None
    formula_id = request.args.get('formula_id', type=int)
    if formula_id:
        formula = conn.execute("SELECT f.*, m.codigo, m.nombre, m.unidad FROM op_formulas f JOIN materiales m "
                               "ON m.id = f.material_id WHERE f.id = ? AND f.estatus = 'VIGENTE'", (formula_id,)).fetchone()
    if formula:
        tamano = request.args.get('tamano_lote', type=float) or formula['tamano_lote']
        presentaciones = srv_formulas.presentaciones(conn, formula['material_id'])
        piezas = {p['id']: request.args.get(f'piezas_{p["id"]}', type=int) or 0 for p in presentaciones}
        disponibilidad = srv_ordenes.disponibilidad(conn, formula['id'], tamano / formula['tamano_lote'], piezas)
        necesario = srv_ordenes.granel_necesario(conn, formula['id'], piezas)
    else:
        piezas, necesario = {}, 0
    conn.close()
    return render_page('ordenes/nueva.html', formulas=formulas, formula=formula, tamano=tamano,
                       presentaciones=presentaciones or [], piezas=piezas, disponibilidad=disponibilidad or [],
                       necesario=necesario, hoy=datetime.now().strftime('%Y-%m-%d'))


@bp.route('/<int:orden_id>')
@login_required
@requiere_modulo('ordenes')
def detalle(orden_id):
    conn = get_db_connection()
    if not srv_ordenes.orden_o_none(conn, orden_id):
        conn.close()
        abort(404)
    exp = srv_expediente.armar(conn, orden_id)
    conn.close()
    return render_page('ordenes/detalle.html', e=exp, o=exp['orden'], nombres=srv_ordenes.NOMBRE_ESTATUS,
                       desviaciones_instalado=srv_ordenes.desviaciones_instalado())


@bp.route('/<int:orden_id>/emitir', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
@requiere_facultad('emitir_ordenes')
def emitir(orden_id):
    def hacer(conn):
        o = srv_ordenes.orden(conn, orden_id)
        fid = firmar(conn, 'firma', 'Autorizó', 'op_ordenes', orden_id, f'Emitió {o["folio"]}',
                     facultad='emitir_ordenes', critica=True)
        return srv_ordenes.emitir(conn, orden_id, fid)
    return accion(hacer, lambda lote: f'Orden emitida con el lote {lote}.', url_for('ordenes.detalle', orden_id=orden_id))


@bp.route('/<int:orden_id>/cancelar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def cancelar(orden_id):
    def hacer(conn):
        o = srv_ordenes.orden(conn, orden_id)
        fid = None
        if o['estatus'] != 'BORRADOR':
            fid = firmar(conn, 'firma', 'Anuló', 'op_ordenes', orden_id, f'Canceló {o["folio"]}',
                         facultad='emitir_ordenes', critica=True)
        srv_ordenes.cancelar(conn, orden_id, motivo=request.form.get('motivo'), usuario=usuario(), firma_id=fid)
    return accion(hacer, 'Orden cancelada.', url_for('ordenes.detalle', orden_id=orden_id))


@bp.route('/<int:orden_id>/revisar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def revisar(orden_id):
    def hacer(conn):
        o = srv_ordenes.orden(conn, orden_id)
        termino = conn.execute('SELECT numero_empleado FROM firmas WHERE id = ?', (o['termino_firma_id'],)).fetchone()
        fid = firmar(conn, 'firma', 'Revisó', 'op_ordenes', orden_id, f'Revisó el expediente {o["folio"]}',
                     roles=ROLES_SUPERVISOR, critica=True,
                     distinto_de=termino['numero_empleado'] if termino else None)
        srv_ordenes.revisar(conn, orden_id, fid)
    return accion(hacer, 'Expediente revisado. Ya lo puede liberar Calidad.', url_for('ordenes.detalle', orden_id=orden_id))


@bp.route('/<int:orden_id>/expediente')
@login_required
@requiere_modulo('ordenes')
def expediente(orden_id):
    from nucleo.parametros import parametro
    conn = get_db_connection()
    if not srv_ordenes.orden_o_none(conn, orden_id):
        conn.close()
        abort(404)
    exp = srv_expediente.armar(conn, orden_id)
    codigo = parametro('ordenes.codigo_formato_expediente', conn)
    conn.close()
    return full_page('ordenes/expediente.html', e=exp, o=exp['orden'], codigo_formato=codigo,
                     nombres=srv_ordenes.NOMBRE_ESTATUS)


@bp.route('/<int:orden_id>/orden-impresa')
@login_required
@requiere_modulo('ordenes')
def orden_impresa(orden_id):
    conn = get_db_connection()
    o = srv_ordenes.orden_o_none(conn, orden_id)
    if not o:
        conn.close()
        abort(404)
    e = srv_ordenes.estado(conn, orden_id)
    if o['estatus'] == 'BORRADOR':
        piezas = {p['presentacion_id']: p['piezas_planeadas'] for p in e['presentaciones']}
        comps = srv_ordenes.requerimientos(conn, o['formula_id'], o['factor'], piezas)
        pasos = srv_formulas.pasos(conn, o['formula_id'])
    else:
        comps = [{**c, 'cantidad': c['cantidad_teorica']} for c in e['componentes']]
        pasos = [{**p, 'parametros': p['valores']} for p in e['pasos']]
    conn.close()
    return full_page('ordenes/orden_impresa.html', o=o, e=e, comps=comps, pasos=pasos)


# ---------------------------------------------------------------------------
# Fórmulas maestras
# ---------------------------------------------------------------------------

@bp.route('/formulas')
@login_required
@requiere_modulo('ordenes')
def formulas():
    conn = get_db_connection()
    productos = []
    for g in conn.execute("SELECT * FROM materiales WHERE tipo = 'GR' AND activo = 1 ORDER BY codigo").fetchall():
        versiones = conn.execute("SELECT * FROM op_formulas WHERE material_id = ? AND estatus <> 'OBSOLETA' "
                                 'ORDER BY version DESC', (g['id'],)).fetchall()
        productos.append({'granel': g, 'versiones': versiones,
                          'presentaciones': srv_formulas.presentaciones(conn, g['id'])})
    conn.close()
    return render_page('ordenes/formulas.html', productos=productos)


@bp.route('/formulas/producto', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_producto():
    f = request.form
    presentaciones = []
    for i in range(1, 6):
        if (f.get(f'p{i}_codigo') or '').strip():
            presentaciones.append({'codigo': f.get(f'p{i}_codigo'), 'nombre': f.get(f'p{i}_nombre'),
                                   'unidad': f.get(f'p{i}_unidad') or 'pza',
                                   'contenido_por_pieza': f.get(f'p{i}_contenido')})

    def hacer(conn):
        gid = srv_formulas.crear_producto(conn, granel={
            'codigo': f.get('codigo'), 'nombre': f.get('nombre'), 'unidad': f.get('unidad') or 'kg',
            'vida_util_meses': f.get('vida_util_meses') or None, 'registro_sanitario': f.get('registro_sanitario'),
            'forma_farmaceutica': f.get('forma_farmaceutica')}, presentaciones=presentaciones, usuario=usuario())
        return srv_formulas.crear_formula(conn, material_id=gid, tamano_lote=f.get('tamano_lote'), usuario=usuario(),
                                          codigo_documento=f.get('codigo_documento'))
    return accion(hacer, 'Producto creado con su fórmula en borrador.',
                  lambda fid: url_for('ordenes.formula_editor', formula_id=fid), url_for('ordenes.formulas'))


@bp.route('/formulas/nueva', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_nueva():
    def hacer(conn):
        return srv_formulas.crear_formula(conn, material_id=request.form.get('material_id', type=int),
                                          tamano_lote=request.form.get('tamano_lote'), usuario=usuario(),
                                          codigo_documento=request.form.get('codigo_documento'))
    return accion(hacer, 'Fórmula creada en borrador.', lambda fid: url_for('ordenes.formula_editor', formula_id=fid),
                  url_for('ordenes.formulas'))


@bp.route('/formulas/material', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_material():
    """Alta rápida de una materia prima o material de envase/empaque desde el editor."""
    from nucleo import inventario
    volver = request.form.get('volver') or url_for('ordenes.formulas')
    if not volver.startswith('/ordenes/') or volver.startswith('//'):
        volver = url_for('ordenes.formulas')
    f = request.form

    def hacer(conn):
        tipo = f.get('tipo')
        if tipo not in ('MP', 'ME', 'EM'):
            raise ErrorOrdenes('Tipo de material no válido.')
        return inventario.crear_material(conn, codigo=f.get('codigo'), nombre=f.get('nombre'), tipo=tipo,
                                         unidad=f.get('unidad') or ('kg' if tipo == 'MP' else 'pza'),
                                         usuario=usuario(), es_impreso=booleano(f.get('es_impreso')),
                                         requiere_analisis=1)
    return accion(hacer, 'Material dado de alta.', volver)


@bp.route('/formulas/<int:formula_id>')
@login_required
@requiere_modulo('ordenes')
def formula_editor(formula_id):
    conn = get_db_connection()
    try:
        f = srv_formulas.formula(conn, formula_id)
    except ErrorOrdenes:
        conn.close()
        abort(404)
    comps = srv_formulas.componentes(conn, formula_id)
    total = sum(c['cantidad'] for c in comps)
    firmas = {r['id']: dict(r) for r in conn.execute('SELECT * FROM firmas WHERE id IN (?, ?)',
                                                     (f['reviso_firma_id'] or 0, f['aprobo_firma_id'] or 0)).fetchall()}
    contexto = dict(
        f=f, componentes=comps, total=total, pasos=srv_formulas.pasos(conn, formula_id),
        empaque=srv_formulas.empaque(conn, formula_id), presentaciones=srv_formulas.presentaciones(conn, f['material_id']),
        materias=conn.execute("SELECT id, codigo, nombre, unidad FROM materiales WHERE tipo = 'MP' AND activo = 1 "
                              'ORDER BY codigo').fetchall(),
        envases=conn.execute("SELECT id, codigo, nombre, unidad, tipo FROM materiales WHERE tipo IN ('ME','EM') "
                             'AND activo = 1 ORDER BY codigo').fetchall(),
        procesos=conn.execute('SELECT id, nombre FROM recetas ORDER BY nombre').fetchall(),
        tareas=conn.execute('SELECT id, nombre FROM competencias_tareas WHERE activo = 1 ORDER BY orden, nombre').fetchall(),
        tipos_equipo=[r[0] for r in conn.execute('SELECT DISTINCT tipo FROM equipos WHERE tipo IS NOT NULL ORDER BY tipo')],
        versiones=conn.execute('SELECT id, version, estatus FROM op_formulas WHERE material_id = ? ORDER BY version DESC',
                               (f['material_id'],)).fetchall(),
        firmas=firmas, pestana=request.args.get('pestana', 'datos'),
        editable=f['estatus'] == 'BORRADOR' and usuario_puede('ordenes_planear'))
    conn.close()
    return render_page('ordenes/formula_editor.html', **contexto)


def _volver(formula_id, pestana='datos'):
    return url_for('ordenes.formula_editor', formula_id=formula_id, pestana=pestana)


@bp.route('/formulas/<int:formula_id>/datos', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_datos(formula_id):
    f = request.form
    return accion(lambda conn: srv_formulas.editar_formula(
        conn, formula_id, tamano_lote=f.get('tamano_lote'), rendimiento_min=f.get('rendimiento_min'),
        rendimiento_max=f.get('rendimiento_max'), codigo_documento=f.get('codigo_documento'), cambios=f.get('cambios')),
        'Datos guardados.', _volver(formula_id))


@bp.route('/formulas/<int:formula_id>/componente', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_componente(formula_id):
    f = request.form
    return accion(lambda conn: srv_formulas.agregar_componente(
        conn, formula_id, material_id=f.get('material_id', type=int), cantidad=f.get('cantidad'),
        tolerancia_pct=f.get('tolerancia_pct') or 1, es_activo=f.get('es_activo'), notas=f.get('notas')),
        'Materia prima agregada.', _volver(formula_id, 'materias'))


@bp.route('/formulas/<int:formula_id>/componente/<int:item_id>/quitar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_componente_quitar(formula_id, item_id):
    return accion(lambda conn: srv_formulas.quitar_componente(conn, formula_id, item_id), 'Materia prima quitada.',
                  _volver(formula_id, 'materias'))


@bp.route('/formulas/<int:formula_id>/paso', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_paso(formula_id):
    f = request.form
    parametros = [{'nombre': n, 'unidad': u, 'valor_min': a, 'valor_max': b}
                  for n, u, a, b in zip(f.getlist('param_nombre'), f.getlist('param_unidad'),
                                        f.getlist('param_min'), f.getlist('param_max'))]
    return accion(lambda conn: srv_formulas.agregar_paso(
        conn, formula_id, instruccion=f.get('instruccion'), fase=f.get('fase') or 'Fabricación',
        minutos=f.get('minutos') or 0, es_critico=f.get('es_critico'), equipo_tipo=f.get('equipo_tipo'),
        competencia_tarea_id=f.get('competencia_tarea_id', type=int), parametros=parametros),
        'Paso agregado.', _volver(formula_id, 'pasos'))


@bp.route('/formulas/<int:formula_id>/paso/<int:item_id>/quitar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_paso_quitar(formula_id, item_id):
    return accion(lambda conn: srv_formulas.quitar_paso(conn, formula_id, item_id), 'Paso quitado.',
                  _volver(formula_id, 'pasos'))


@bp.route('/formulas/<int:formula_id>/importar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_importar(formula_id):
    return accion(lambda conn: srv_formulas.importar_de_proceso(
        conn, formula_id, request.form.get('receta_id', type=int), current_app.static_folder,
        fase=request.form.get('fase') or 'Fabricación'),
        lambda n: f'Se importaron {n} pasos.', _volver(formula_id, 'pasos'))


@bp.route('/formulas/<int:formula_id>/empaque', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_empaque(formula_id):
    f = request.form
    return accion(lambda conn: srv_formulas.agregar_empaque(
        conn, formula_id, presentacion_id=f.get('presentacion_id', type=int), material_id=f.get('material_id', type=int),
        cantidad_por_pieza=f.get('cantidad_por_pieza'), tolerancia_pct=f.get('tolerancia_pct') or 2),
        'Material agregado.', _volver(formula_id, 'empaque'))


@bp.route('/formulas/<int:formula_id>/empaque/<int:item_id>/quitar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_empaque_quitar(formula_id, item_id):
    return accion(lambda conn: srv_formulas.quitar_empaque(conn, formula_id, item_id), 'Material quitado.',
                  _volver(formula_id, 'empaque'))


@bp.route('/formulas/<int:formula_id>/presentacion', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_presentacion(formula_id):
    f = request.form

    def hacer(conn):
        formula = srv_formulas.formula(conn, formula_id)
        return srv_formulas.agregar_presentacion(conn, formula['material_id'], {
            'codigo': f.get('codigo'), 'nombre': f.get('nombre'), 'unidad': f.get('unidad') or 'pza',
            'contenido_por_pieza': f.get('contenido_por_pieza')}, usuario=usuario())
    return accion(hacer, 'Presentación agregada.', _volver(formula_id, 'empaque'))


@bp.route('/formulas/<int:formula_id>/enviar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_enviar(formula_id):
    return accion(lambda conn: srv_formulas.enviar_revision(conn, formula_id), 'Fórmula enviada a revisión.',
                  _volver(formula_id))


@bp.route('/formulas/<int:formula_id>/regresar', methods=['POST'])
@login_required
@requiere_modulo('ordenes')
def formula_regresar(formula_id):
    if not (usuario_puede('ordenes_planear') or 'aprobar_maestros' in (session.get('facultades') or [])):
        flash('No tienes permiso para regresar la fórmula.', 'error')
        return hx_redirect(_volver(formula_id))
    return accion(lambda conn: srv_formulas.regresar_a_borrador(conn, formula_id), 'La fórmula regresó a borrador.',
                  _volver(formula_id))


@bp.route('/formulas/<int:formula_id>/revisar', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_revisar(formula_id):
    def hacer(conn):
        f = srv_formulas.formula(conn, formula_id)
        fid = firmar(conn, 'firma', 'Revisó', 'op_formulas', formula_id,
                     f'Revisó la fórmula {f["material_codigo"]} v{f["version"]}', roles=ROLES_SUPERVISOR, critica=True)
        srv_formulas.revisar(conn, formula_id, fid)
    return accion(hacer, 'Fórmula revisada. Falta la aprobación de Calidad.', _volver(formula_id))


@bp.route('/formulas/<int:formula_id>/aprobar', methods=['POST'])
@login_required
@requiere_modulo('ordenes')
@requiere_facultad('aprobar_maestros')
def formula_aprobar(formula_id):
    def hacer(conn):
        f = srv_formulas.formula(conn, formula_id)
        revisor = conn.execute('SELECT numero_empleado FROM firmas WHERE id = ?', (f['reviso_firma_id'] or 0,)).fetchone()
        fid = firmar(conn, 'firma', 'Aprobó', 'op_formulas', formula_id,
                     f'Aprobó la fórmula {f["material_codigo"]} v{f["version"]}', facultad='aprobar_maestros',
                     critica=True, distinto_de=revisor['numero_empleado'] if revisor else None)
        srv_formulas.aprobar(conn, formula_id, fid)
    return accion(hacer, 'Fórmula aprobada: ya está vigente.', _volver(formula_id))


@bp.route('/formulas/<int:formula_id>/nueva-version', methods=['POST'])
@login_required
@requiere_modulo('ordenes_planear')
def formula_nueva_version(formula_id):
    return accion(lambda conn: srv_formulas.nueva_version(conn, formula_id, usuario=usuario(),
                                                          cambios=request.form.get('cambios')),
                  'Nueva versión creada en borrador.', lambda fid: _volver(fid), _volver(formula_id))


# ---------------------------------------------------------------------------
# Reporte en Excel
# ---------------------------------------------------------------------------

@bp.route('/reporte')
@login_required
@requiere_modulo('ordenes')
def reporte():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    conn = get_db_connection()
    wb = Workbook()
    hoja = wb.active
    hoja.title = 'Rendimiento'
    hoja.append(['Producto', 'Mes', 'Lotes', 'Rendimiento promedio %', 'Mínimo %', 'Máximo %'])
    for r in conn.execute("SELECT m.codigo || ' ' || m.nombre AS producto, substr(o.fecha_fin, 1, 7) AS mes, COUNT(*) AS n, "
                          'ROUND(AVG(o.rendimiento_pct), 2), MIN(o.rendimiento_pct), MAX(o.rendimiento_pct) '
                          'FROM op_ordenes o JOIN materiales m ON m.id = o.material_id '
                          "WHERE o.estatus IN ('TERMINADA','REVISADA') GROUP BY o.material_id, mes "
                          'ORDER BY mes DESC, producto').fetchall():
        hoja.append(list(r))
    hoja = wb.create_sheet('Por estatus')
    hoja.append(['Estatus', 'Órdenes'])
    for r in conn.execute('SELECT estatus, COUNT(*) FROM op_ordenes GROUP BY estatus ORDER BY estatus').fetchall():
        hoja.append([srv_ordenes.NOMBRE_ESTATUS.get(r[0], r[0]), r[1]])
    hoja = wb.create_sheet('Tiempos')
    hoja.append(['Folio', 'Producto', 'Lote', 'Estatus', 'Emitida', 'Inicio', 'Terminada', 'Revisada', 'Liberada'])
    for r in conn.execute(
            "SELECT o.folio, m.codigo || ' ' || m.nombre, o.numero_lote, o.estatus, o.fecha_emision, o.fecha_inicio, "
            'o.fecha_fin, o.fecha_revision, (SELECT MAX(h.ts) FROM lote_historial h WHERE h.lote_id = o.granel_lote_id '
            "AND h.campo = 'estatus' AND h.valor_nuevo = 'APROBADO') FROM op_ordenes o "
            'JOIN materiales m ON m.id = o.material_id ORDER BY o.id DESC LIMIT 5000').fetchall():
        fila = list(r)
        fila[3] = srv_ordenes.NOMBRE_ESTATUS.get(fila[3], fila[3])
        hoja.append(fila)
    conn.close()
    for h in wb.worksheets:
        for celda in h[1]:
            celda.font = Font(bold=True)
        for col in h.columns:
            h.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or '')) for c in col) + 2))
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return send_file(salida, as_attachment=True, download_name=f'ordenes_produccion_{datetime.now():%Y%m%d}.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
