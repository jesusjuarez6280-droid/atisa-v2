"""Reportes de Calidad en Excel (diseño 03, 4.8)."""
import io

from .servicios._comun import fecha, texto

REPORTES = {
    'analisis': 'Análisis por periodo con sus tiempos',
    'proveedores': 'Lotes rechazados e investigaciones por proveedor',
    'investigaciones': 'Investigaciones por prueba y material',
    'tendencia': 'Resultados de una prueba lote por lote',
}


def _filas_analisis(conn, desde, hasta, **_):
    titulos = ['Folio', 'Tipo', 'Estatus', 'Material', 'Nombre', 'Lote', 'Entrada a cuarentena', 'Muestreo',
               'Último resultado', 'Terminado', 'Dictamen', 'Decisión', 'Días cuarentena→muestreo',
               'Días muestreo→resultados', 'Días resultados→dictamen', 'Días totales']
    filas = conn.execute(
        'SELECT a.folio, a.tipo, a.estatus, m.codigo, m.nombre, l.numero_lote, l.fecha_entrada, a.creado_en, '
        '(SELECT MAX(r.ts) FROM cc_resultados r WHERE r.analisis_id = a.id) AS ultimo, a.terminado_en, '
        '(SELECT d.ts FROM cc_dictamenes d WHERE d.analisis_id = a.id ORDER BY d.id LIMIT 1) AS dict_ts, '
        '(SELECT d.decision FROM cc_dictamenes d WHERE d.analisis_id = a.id ORDER BY d.id LIMIT 1) AS decision '
        'FROM cc_analisis a JOIN lotes l ON l.id = a.lote_id JOIN materiales m ON m.id = l.material_id '
        'WHERE date(a.creado_en) BETWEEN ? AND ? ORDER BY a.id', (desde, hasta)).fetchall()

    def dias(a, b):
        if not a or not b:
            return None
        r = conn.execute('SELECT ROUND(julianday(?) - julianday(?), 1)', (b, a)).fetchone()[0]
        return r
    salida = []
    for f in filas:
        salida.append([f['folio'], f['tipo'], f['estatus'], f['codigo'], f['nombre'], f['numero_lote'],
                       f['fecha_entrada'], f['creado_en'], f['ultimo'], f['terminado_en'], f['dict_ts'], f['decision'],
                       dias(f['fecha_entrada'], f['creado_en']), dias(f['creado_en'], f['ultimo']),
                       dias(f['ultimo'], f['dict_ts']), dias(f['fecha_entrada'], f['dict_ts'])])
    return titulos, salida


def _filas_proveedores(conn, desde, hasta, **_):
    titulos = ['Proveedor', 'Estatus del proveedor', 'Lotes recibidos', 'Lotes aprobados', 'Lotes rechazados',
               'Investigaciones OOS', 'OOS confirmadas']
    filas = conn.execute(
        'SELECT p.nombre, p.estatus, '
        'COUNT(DISTINCT l.id) AS recibidos, '
        "COUNT(DISTINCT CASE WHEN d.decision = 'APROBADO' THEN l.id END) AS aprobados, "
        "COUNT(DISTINCT CASE WHEN d.decision = 'RECHAZADO' THEN l.id END) AS rechazados, "
        'COUNT(DISTINCT i.id) AS oos, '
        "COUNT(DISTINCT CASE WHEN i.conclusion = 'CONFIRMADO' THEN i.id END) AS confirmadas "
        'FROM proveedores p JOIN lotes l ON l.proveedor_id = p.id '
        'LEFT JOIN cc_dictamenes d ON d.lote_id = l.id '
        'LEFT JOIN cc_analisis a ON a.lote_id = l.id LEFT JOIN cc_investigaciones i ON i.analisis_id = a.id '
        'WHERE l.fecha_entrada BETWEEN ? AND ? GROUP BY p.id ORDER BY rechazados DESC, oos DESC, p.nombre',
        (desde, hasta)).fetchall()
    return titulos, [list(f) for f in filas]


def _filas_investigaciones(conn, desde, hasta, **_):
    titulos = ['Folio', 'Estatus', 'Conclusión', 'Análisis', 'Material', 'Nombre', 'Lote', 'Prueba', 'Resultado',
               'Especificación', 'Hallazgos', 'Abierta', 'Cerrada']
    filas = conn.execute(
        'SELECT i.folio, i.estatus, i.conclusion, a.folio, m.codigo, m.nombre, l.numero_lote, p.prueba, '
        'COALESCE(CAST(r.valor_redondeado AS TEXT), r.texto), p.texto_especificacion, i.hallazgos, i.abierta_en, '
        'i.cerrada_en FROM cc_investigaciones i JOIN cc_analisis a ON a.id = i.analisis_id '
        'JOIN lotes l ON l.id = a.lote_id JOIN materiales m ON m.id = l.material_id '
        'JOIN cc_resultados r ON r.id = i.resultado_id JOIN cc_especificacion_pruebas p ON p.id = r.prueba_id '
        'WHERE date(i.abierta_en) BETWEEN ? AND ? ORDER BY p.prueba, m.codigo, i.id', (desde, hasta)).fetchall()
    return titulos, [list(f) for f in filas]


def _filas_tendencia(conn, desde, hasta, material=None, prueba=None, **_):
    titulos = ['Material', 'Lote', 'Análisis', 'Prueba', 'Resultado', 'Unidad', 'Límite mínimo', 'Límite máximo',
               'Cumple', 'Válido', 'Fecha del análisis', 'Fuente']
    sql = ('SELECT m.codigo, l.numero_lote, a.folio, p.prueba, COALESCE(CAST(r.valor_redondeado AS TEXT), r.texto), '
           "p.unidad, p.limite_min, p.limite_max, CASE r.cumple WHEN 1 THEN 'Sí' ELSE 'No' END, "
           "CASE r.valido WHEN 1 THEN 'Sí' ELSE 'Anulado' END, r.fecha_analisis, r.fuente "
           'FROM cc_resultados r JOIN cc_analisis a ON a.id = r.analisis_id JOIN lotes l ON l.id = a.lote_id '
           'JOIN materiales m ON m.id = l.material_id JOIN cc_especificacion_pruebas p ON p.id = r.prueba_id '
           'WHERE r.fecha_analisis BETWEEN ? AND ? ')
    params = [desde, hasta]
    if texto(material):
        sql += 'AND (m.codigo = ? OR m.nombre LIKE ?) '
        params += [material.strip(), f'%{material.strip()}%']
    if texto(prueba):
        sql += 'AND p.prueba LIKE ? '
        params.append(f'%{prueba.strip()}%')
    return titulos, [list(f) for f in conn.execute(sql + 'ORDER BY m.codigo, p.prueba, r.fecha_analisis, r.id',
                                                   params).fetchall()]


_ARMADORES = {'analisis': _filas_analisis, 'proveedores': _filas_proveedores,
              'investigaciones': _filas_investigaciones, 'tendencia': _filas_tendencia}


def excel(conn, tipo, desde=None, hasta=None, **filtros) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    desde = fecha(desde, 'fecha desde') or '2000-01-01'
    hasta = fecha(hasta, 'fecha hasta') or '2999-12-31'
    titulos, filas = _ARMADORES[tipo](conn, desde, hasta, **filtros)
    wb = Workbook()
    ws = wb.active
    ws.title = REPORTES[tipo][:31]
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append(list(f))
    for i, t in enumerate(titulos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(12, min(40, len(t) + 4))
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
