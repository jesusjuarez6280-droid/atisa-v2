"""Especificaciones con versiones: qué pruebas, con qué método y con qué límites (diseño 03, 4.6 y regla 1).

BORRADOR → EN_REVISION (Revisó) → VIGENTE (Aprobó, con aprobar_maestros y distinto de quien revisó).
Una sola vigente por material; la vigente no se edita: se hace versión nueva y la anterior queda
OBSOLETA al aprobarse. Cada análisis guarda la versión con la que empezó.
"""
import io

from nucleo import firmas

from ._comun import (ErrorCalidad, ROLES_CALIDAD, texto, numero, booleano, hoy, firmar,
                     numero_firmante)

TIPOS_PRUEBA = ('NUMERICA', 'CUALITATIVA')
FUENTES = ('Interno', 'Externo', 'Certificado del proveedor')
CAMPOS = ('codigo_documento', 'referencia', 'documento_id', 'cantidad_muestra', 'plan_muestreo',
          'pide_retencion', 'cantidad_retencion', 'meses_reanalisis', 'requiere_revision', 'cambios')
CAMPOS_PRUEBA = ('prueba', 'tipo', 'limite_min', 'limite_max', 'unidad', 'decimales', 'texto_especificacion',
                 'metodo', 'fuente_default', 'equipo_tipo', 'cada_lote', 'en_certificado')
# Envase y empaque: el dictamen hace de segunda revisión, así que se puede quitar la revisión.
TIPOS_SIN_REVISION = ('ME', 'EM')


def especificacion(conn, especificacion_id):
    fila = conn.execute(
        'SELECT e.*, m.codigo AS material_codigo, m.nombre AS material_nombre, m.tipo AS material_tipo, '
        'm.unidad AS unidad, m.es_impreso FROM cc_especificaciones e JOIN materiales m ON m.id = e.material_id '
        'WHERE e.id = ?', (especificacion_id,)).fetchone()
    if not fila:
        raise ErrorCalidad('La especificación no existe.')
    return fila


def pruebas(conn, especificacion_id):
    return conn.execute('SELECT * FROM cc_especificacion_pruebas WHERE especificacion_id = ? ORDER BY orden, id',
                        (especificacion_id,)).fetchall()


def vigente(conn, material_id):
    return conn.execute("SELECT * FROM cc_especificaciones WHERE material_id = ? AND estatus = 'VIGENTE'",
                        (material_id,)).fetchone()


def _en_curso(conn, material_id):
    return conn.execute("SELECT * FROM cc_especificaciones WHERE material_id = ? "
                        "AND estatus IN ('BORRADOR','EN_REVISION') ORDER BY version DESC LIMIT 1",
                        (material_id,)).fetchone()


def _fmt(n, decimales):
    return f'{n:.{int(decimales)}f}'


def texto_automatico(p):
    """'98.0 – 102.0 %', 'Mín. 98.0 %', 'Máx. 0.5 %' a partir de los límites."""
    d, u = p.get('decimales') or 0, (' ' + p['unidad']) if p.get('unidad') else ''
    lo, hi = p.get('limite_min'), p.get('limite_max')
    if lo is not None and hi is not None:
        return f'{_fmt(lo, d)} – {_fmt(hi, d)}{u}'
    if lo is not None:
        return f'Mín. {_fmt(lo, d)}{u}'
    return f'Máx. {_fmt(hi, d)}{u}'


def _validar_prueba(p, n):
    p = {k: p.get(k) for k in CAMPOS_PRUEBA}
    p['prueba'] = texto(p['prueba'])
    if not p['prueba']:
        raise ErrorCalidad(f'Prueba {n}: falta el nombre.')
    nombre = p['prueba']
    p['tipo'] = (texto(p['tipo']) or 'NUMERICA').upper().replace('Ú', 'U').replace('É', 'E')
    if p['tipo'] not in TIPOS_PRUEBA:
        raise ErrorCalidad(f'{nombre}: el tipo debe ser numérica o cualitativa.')
    p['decimales'] = numero(p['decimales'], 'los decimales', entero=True)
    if p['decimales'] is None:
        p['decimales'] = 1
    if not 0 <= p['decimales'] <= 6:
        raise ErrorCalidad(f'{nombre}: los decimales van de 0 a 6.')
    if p['tipo'] == 'NUMERICA':
        p['limite_min'] = numero(p['limite_min'], 'el límite mínimo')
        p['limite_max'] = numero(p['limite_max'], 'el límite máximo')
        if p['limite_min'] is None and p['limite_max'] is None:
            raise ErrorCalidad(f'{nombre}: una prueba numérica necesita al menos un límite.')
        if p['limite_min'] is not None and p['limite_max'] is not None and p['limite_min'] > p['limite_max']:
            raise ErrorCalidad(f'{nombre}: el límite mínimo es mayor que el máximo.')
    else:
        p['limite_min'] = p['limite_max'] = None
    p['unidad'] = texto(p['unidad'])
    p['texto_especificacion'] = texto(p['texto_especificacion'])
    if not p['texto_especificacion']:
        if p['tipo'] == 'CUALITATIVA':
            raise ErrorCalidad(f'{nombre}: escribe la especificación (lo que debe observarse).')
        p['texto_especificacion'] = texto_automatico(p)
    p['metodo'] = texto(p['metodo'])
    p['fuente_default'] = texto(p['fuente_default']) or 'Interno'
    if p['fuente_default'] not in FUENTES:
        raise ErrorCalidad(f'{nombre}: la fuente debe ser Interno, Externo o Certificado del proveedor.')
    p['equipo_tipo'] = texto(p['equipo_tipo'])
    p['cada_lote'] = 1 if p['cada_lote'] is None else booleano(p['cada_lote'])
    p['en_certificado'] = 1 if p['en_certificado'] is None else booleano(p['en_certificado'])
    return p


def _validar_datos(conn, datos, material):
    datos = {k: v for k, v in (datos or {}).items() if k in CAMPOS}
    for campo in ('codigo_documento', 'referencia', 'plan_muestreo', 'cambios'):
        if campo in datos:
            datos[campo] = texto(datos[campo])
    for campo, nombre in (('cantidad_muestra', 'la cantidad a muestrear'),
                          ('cantidad_retencion', 'la cantidad de retención')):
        if campo in datos:
            datos[campo] = numero(datos[campo], nombre)
            if datos[campo] is not None and datos[campo] < 0:
                raise ErrorCalidad(f'{nombre.capitalize()} no puede ser negativa.')
    if 'meses_reanalisis' in datos:
        datos['meses_reanalisis'] = numero(datos['meses_reanalisis'], 'los meses de reanálisis', entero=True)
        if datos['meses_reanalisis'] is not None and datos['meses_reanalisis'] <= 0:
            datos['meses_reanalisis'] = None
    if 'documento_id' in datos:
        datos['documento_id'] = numero(datos['documento_id'], 'el documento', entero=True) or None
        if datos['documento_id'] and not conn.execute('SELECT 1 FROM documentos_legales WHERE id = ?',
                                                      (datos['documento_id'],)).fetchone():
            raise ErrorCalidad('El documento ligado no existe en Documentos Legales.')
    for campo in ('pide_retencion', 'requiere_revision'):
        if campo in datos:
            datos[campo] = booleano(datos[campo])
    if datos.get('requiere_revision') == 0 and material['tipo'] not in TIPOS_SIN_REVISION:
        raise ErrorCalidad('Solo envase y empaque pueden quitar la revisión de una segunda persona.')
    if datos.get('pide_retencion') and not datos.get('cantidad_retencion'):
        raise ErrorCalidad('Indica cuánto se guarda de muestra de retención.')
    return datos


def _material(conn, material_id):
    m = conn.execute('SELECT * FROM materiales WHERE id = ?', (material_id,)).fetchone()
    if not m:
        raise ErrorCalidad('El material no existe.')
    return m


def _insertar_pruebas(conn, especificacion_id, lista):
    conn.execute('DELETE FROM cc_especificacion_pruebas WHERE especificacion_id = ?', (especificacion_id,))
    for i, p in enumerate(lista, start=1):
        conn.execute(f'INSERT INTO cc_especificacion_pruebas (especificacion_id, orden, {", ".join(CAMPOS_PRUEBA)}) '
                     f'VALUES (?, ?, {", ".join("?" * len(CAMPOS_PRUEBA))})',
                     (especificacion_id, i, *(p[k] for k in CAMPOS_PRUEBA)))


def crear(conn, *, material_id, usuario, datos=None, pruebas_lista=None) -> int:
    """Borrador nuevo (versión siguiente). Solo un borrador o revisión a la vez por material."""
    material = _material(conn, material_id)
    otra = _en_curso(conn, material_id)
    if otra:
        raise ErrorCalidad(f'{material["codigo"]} ya tiene la versión {otra["version"]} en '
                           f'{"borrador" if otra["estatus"] == "BORRADOR" else "revisión"}; termínala primero.')
    datos = _validar_datos(conn, datos, material)
    lista = [_validar_prueba(p, i) for i, p in enumerate(pruebas_lista or [], start=1)]
    version = conn.execute('SELECT COALESCE(MAX(version), 0) + 1 FROM cc_especificaciones WHERE material_id = ?',
                           (material_id,)).fetchone()[0]
    cols = {'material_id': material_id, 'version': version, 'elaboro': usuario, **datos}
    cur = conn.execute(f'INSERT INTO cc_especificaciones ({", ".join(cols)}) VALUES ({", ".join("?" * len(cols))})',
                       tuple(cols.values()))
    _insertar_pruebas(conn, cur.lastrowid, lista)
    return cur.lastrowid


def guardar(conn, especificacion_id, *, usuario, datos=None, pruebas_lista=None):
    """Cambia datos y pruebas de un borrador (las pruebas se reemplazan completas si vienen)."""
    e = especificacion(conn, especificacion_id)
    if e['estatus'] != 'BORRADOR':
        raise ErrorCalidad('Solo un borrador se edita; una vigente cambia con una versión nueva.')
    datos = _validar_datos(conn, datos, _material(conn, e['material_id']))
    if datos:
        conn.execute(f'UPDATE cc_especificaciones SET {", ".join(f"{k} = ?" for k in datos)} WHERE id = ?',
                     (*datos.values(), especificacion_id))
    if pruebas_lista is not None:
        _insertar_pruebas(conn, especificacion_id, [_validar_prueba(p, i) for i, p in enumerate(pruebas_lista, 1)])


def _copia(conn, origen):
    datos = {k: origen[k] for k in CAMPOS if k != 'cambios'}
    lista = [{k: p[k] for k in CAMPOS_PRUEBA} for p in pruebas(conn, origen['id'])]
    return datos, lista


def copiar(conn, *, desde_id, material_id, usuario) -> int:
    """Borrador para `material_id` con los datos y pruebas de otra especificación."""
    origen = especificacion(conn, desde_id)
    datos, lista = _copia(conn, origen)
    material = _material(conn, material_id)
    if material['tipo'] not in TIPOS_SIN_REVISION:
        datos['requiere_revision'] = 1
    return crear(conn, material_id=material_id, usuario=usuario, datos=datos, pruebas_lista=lista)


def nueva_version(conn, especificacion_id, *, usuario) -> int:
    origen = especificacion(conn, especificacion_id)
    if origen['estatus'] not in ('VIGENTE', 'OBSOLETA'):
        raise ErrorCalidad('La versión nueva se hace a partir de una especificación aprobada.')
    datos, lista = _copia(conn, origen)
    return crear(conn, material_id=origen['material_id'], usuario=usuario, datos=datos, pruebas_lista=lista)


def enviar_revision(conn, especificacion_id, *, usuario):
    e = especificacion(conn, especificacion_id)
    if e['estatus'] != 'BORRADOR':
        raise ErrorCalidad('Solo un borrador se envía a revisión.')
    if not pruebas(conn, especificacion_id):
        raise ErrorCalidad('Agrega al menos una prueba antes de enviarla a revisión.')
    if e['version'] > 1 and not e['cambios']:
        raise ErrorCalidad('Escribe qué cambió respecto a la versión anterior.')
    conn.execute("UPDATE cc_especificaciones SET estatus = 'EN_REVISION' WHERE id = ?", (especificacion_id,))


def regresar_borrador(conn, especificacion_id, *, usuario):
    e = especificacion(conn, especificacion_id)
    if e['estatus'] != 'EN_REVISION':
        raise ErrorCalidad('Solo una especificación en revisión regresa a borrador.')
    conn.execute("UPDATE cc_especificaciones SET estatus = 'BORRADOR', reviso_firma_id = NULL WHERE id = ?",
                 (especificacion_id,))


def revisar(conn, especificacion_id, *, firma) -> int:
    e = especificacion(conn, especificacion_id)
    if e['estatus'] != 'EN_REVISION':
        raise ErrorCalidad('Solo se revisa una especificación enviada a revisión.')
    if e['reviso_firma_id']:
        raise ErrorCalidad('Esta versión ya está revisada; falta que la aprueben.')
    fid = firmar(conn, firma, 'Revisó', 'cc_especificaciones', especificacion_id,
                 f'Revisó especificación {e["material_codigo"]} v{e["version"]}', roles=ROLES_CALIDAD)
    conn.execute('UPDATE cc_especificaciones SET reviso_firma_id = ? WHERE id = ?', (fid, especificacion_id))
    return fid


def aprobar(conn, especificacion_id, *, firma) -> int:
    """Queda VIGENTE; la vigente anterior del material pasa a OBSOLETA."""
    e = especificacion(conn, especificacion_id)
    if e['estatus'] != 'EN_REVISION':
        raise ErrorCalidad('Solo se aprueba una especificación en revisión.')
    if not e['reviso_firma_id']:
        raise ErrorCalidad('Falta la firma de quien revisó.')
    fid = firmar(conn, firma, 'Aprobó', 'cc_especificaciones', especificacion_id,
                 f'Aprobó especificación {e["material_codigo"]} v{e["version"]}', facultad='aprobar_maestros',
                 distinto_de=numero_firmante(conn, e['reviso_firma_id']))
    conn.execute("UPDATE cc_especificaciones SET estatus = 'OBSOLETA' WHERE material_id = ? AND estatus = 'VIGENTE'",
                 (e['material_id'],))
    conn.execute("UPDATE cc_especificaciones SET estatus = 'VIGENTE', aprobo_firma_id = ?, vigente_desde = ? "
                 'WHERE id = ?', (fid, hoy(), especificacion_id))
    return fid


def en_bloque(conn, accion, ids, *, firma):
    """Revisar o aprobar varias a la vez: el PIN se escribe una vez y queda una firma por especificación.
    Devuelve [(id, error o None)]; las que fallan no detienen a las demás."""
    fn = {'revisar': revisar, 'aprobar': aprobar}[accion]
    salida = []
    for eid in ids:
        conn.execute('SAVEPOINT cc_bloque')
        try:
            fn(conn, int(eid), firma=firma)
            conn.execute('RELEASE cc_bloque')
            salida.append((int(eid), None))
        except (ErrorCalidad, firmas.FirmaInvalida) as e:
            conn.execute('ROLLBACK TO cc_bloque')
            conn.execute('RELEASE cc_bloque')
            salida.append((int(eid), str(e)))
            if isinstance(e, firmas.FirmaInvalida) and 'PIN' in str(e):
                break       # PIN equivocado: no se sigue intentando (cuenta para el bloqueo)
    return salida


# --- Listas -----------------------------------------------------------------

def lista(conn):
    """Por material: la vigente y las que están en borrador o revisión."""
    return conn.execute(
        'SELECT e.*, m.codigo AS material_codigo, m.nombre AS material_nombre, m.tipo AS material_tipo, '
        '(SELECT COUNT(*) FROM cc_especificacion_pruebas p WHERE p.especificacion_id = e.id) AS n_pruebas '
        'FROM cc_especificaciones e JOIN materiales m ON m.id = e.material_id '
        "WHERE e.estatus <> 'OBSOLETA' ORDER BY m.tipo, m.codigo, e.version").fetchall()


def sin_especificacion(conn):
    """Materiales activos que requieren análisis y no tienen vigente (primero los que tienen lotes esperando)."""
    return conn.execute(
        'SELECT m.*, (SELECT COUNT(*) FROM lotes l WHERE l.material_id = m.id AND l.estatus = \'CUARENTENA\') '
        'AS lotes_esperando, (SELECT e.estatus FROM cc_especificaciones e WHERE e.material_id = m.id '
        "AND e.estatus IN ('BORRADOR','EN_REVISION') LIMIT 1) AS en_curso "
        'FROM materiales m WHERE m.activo = 1 AND m.requiere_analisis = 1 '
        "AND NOT EXISTS (SELECT 1 FROM cc_especificaciones e WHERE e.material_id = m.id AND e.estatus = 'VIGENTE') "
        # Las presentaciones de un granel se analizan con la del granel.
        "AND NOT (m.tipo = 'PT' AND m.material_padre_id IS NOT NULL) "
        'ORDER BY lotes_esperando DESC, m.tipo, m.codigo').fetchall()


# --- Carga desde Excel ------------------------------------------------------

COLUMNAS_EXCEL = [
    ('codigo_material', 'Código del material (como está en el catálogo)'),
    ('codigo_documento', 'Código del documento (ESP-MP-012)'),
    ('referencia', 'FEUM, USP, proveedor, interna'),
    ('cantidad_muestra', 'Cantidad a muestrear, en la unidad del material'),
    ('plan_muestreo', 'Plan de muestreo'),
    ('pide_retencion', 'Sí / No'),
    ('cantidad_retencion', 'Cantidad de retención'),
    ('meses_reanalisis', 'Meses de reanálisis (materias primas)'),
    ('prueba', 'Nombre de la prueba'),
    ('tipo', 'NUMERICA o CUALITATIVA'),
    ('limite_min', 'Límite mínimo (numérica)'),
    ('limite_max', 'Límite máximo (numérica)'),
    ('unidad', 'Unidad (%, pH, mg/g…)'),
    ('decimales', 'Decimales a los que se redondea'),
    ('texto_especificacion', 'Cómo se imprime: 98.0 – 102.0 % / Polvo fino blanco'),
    ('metodo', 'Método (MA-012, FEUM MGA 0991…)'),
    ('fuente_default', 'Interno, Externo o Certificado del proveedor'),
    ('equipo_tipo', 'Tipo de equipo (Balanza analítica, Potenciómetro…)'),
    ('cada_lote', 'Sí = en cada lote / No = periódica'),
    ('en_certificado', 'Sí / No'),
]
_DATOS_EXCEL = ('codigo_documento', 'referencia', 'cantidad_muestra', 'plan_muestreo', 'pide_retencion',
                'cantidad_retencion', 'meses_reanalisis')


def plantilla_excel() -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = 'Especificaciones'
    ws.append([c for c, _ in COLUMNAS_EXCEL])
    for celda in ws[1]:
        celda.font = Font(bold=True)
    ws.append(['MP-0001', 'ESP-MP-001', 'FEUM', 0.1, 'raíz de N + 1', 'Sí', 0.05, 24,
               'Valoración', 'NUMERICA', 98.0, 102.0, '%', 1, '', 'MA-012', 'Interno', 'HPLC', 'Sí', 'Sí'])
    ws.append(['MP-0001', '', '', '', '', '', '', '',
               'Aspecto', 'CUALITATIVA', '', '', '', '', 'Polvo fino blanco', 'Visual', 'Interno', '', 'Sí', 'Sí'])
    ayuda = wb.create_sheet('Instrucciones')
    ayuda.append(['Columna', 'Qué va'])
    for c, d in COLUMNAS_EXCEL:
        ayuda.append([c, d])
    ayuda.append([])
    ayuda.append(['', 'Un renglón por prueba. Los datos de la especificación se toman del primer renglón '
                      'de cada material. Se crean en borrador; después se revisan y aprueban en bloque.'])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def _no_vacio(v):
    return v is not None and str(v).strip() != ''


def importar_excel(conn, archivo, *, usuario):
    """Crea un borrador por material del Excel. Devuelve {'creadas': [(codigo, id)], 'errores': [texto]}.
    Un material con errores no se crea; los demás sí."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:
        raise ErrorCalidad('No se pudo leer el archivo: usa la plantilla de Excel (.xlsx).')
    ws = wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise ErrorCalidad('El archivo está vacío.')
    encabezado = [str(c or '').strip().lower() for c in filas[0]]
    if 'codigo_material' not in encabezado or 'prueba' not in encabezado:
        raise ErrorCalidad('Faltan las columnas codigo_material y prueba (usa la plantilla).')
    por_material, orden = {}, []
    for n, fila in enumerate(filas[1:], start=2):
        r = {encabezado[i]: v for i, v in enumerate(fila) if i < len(encabezado) and encabezado[i]}
        codigo = texto(r.get('codigo_material'))
        if not codigo and not _no_vacio(r.get('prueba')):
            continue
        if not codigo:
            por_material.setdefault(None, []).append(n)
            continue
        if codigo not in por_material:
            por_material[codigo] = []
            orden.append(codigo)
        por_material[codigo].append((n, r))
    creadas, errores = [], []
    if None in por_material:
        errores.append(f'Renglones sin código de material: {", ".join(map(str, por_material[None]))}.')
    for codigo in orden:
        renglones = por_material[codigo]
        m = conn.execute('SELECT id FROM materiales WHERE codigo = ?', (codigo,)).fetchone()
        if not m:
            errores.append(f'{codigo}: no existe en el catálogo de materiales.')
            continue
        primero = renglones[0][1]
        datos = {k: primero.get(k) for k in _DATOS_EXCEL if _no_vacio(primero.get(k))}
        lista_p = [{k: r.get(k) for k in CAMPOS_PRUEBA} for _, r in renglones if _no_vacio(r.get('prueba'))]
        conn.execute('SAVEPOINT cc_excel')
        try:
            eid = crear(conn, material_id=m['id'], usuario=usuario, datos=datos, pruebas_lista=lista_p)
            if not lista_p:
                raise ErrorCalidad('no trae pruebas.')
            conn.execute('RELEASE cc_excel')
            creadas.append((codigo, eid))
        except ErrorCalidad as e:
            conn.execute('ROLLBACK TO cc_excel')
            conn.execute('RELEASE cc_excel')
            errores.append(f'{codigo}: {e}')
    return {'creadas': creadas, 'errores': errores}
