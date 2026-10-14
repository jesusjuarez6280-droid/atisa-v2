"""Módulo de calidad y liberación: las pruebas de la sección 9 del diseño 03."""
import io
from datetime import date, timedelta

import pytest
from werkzeug.datastructures import FileStorage

import erp
from fabricas import crear_usuario, crear_material, crear_lote, crear_equipo, conexion, login, PASSWORD
from nucleo import inventario, firmas, registro
from modulos.calidad.servicios import (analisis as an, especificaciones as esp, investigaciones as inv,
                                       dictamen as dic, retencion as ret, certificado as cert)
from modulos.calidad.servicios._comun import ErrorCalidad, redondear, sumar_meses

PIN = '2580'


def F(numero):
    """Firma con PIN fuera de una petición (la sesión es la de quien firma)."""
    return {'numero_empleado': numero, 'pin': PIN,
            'sesion': {'username': 'pruebas', 'numero_empleado': numero, 'ip': None}}


@pytest.fixture(scope='module')
def gente():
    c = conexion()
    g = {
        'analista': crear_usuario(c, 'cc_analista', rol='calidad', pin=PIN),
        'analista2': crear_usuario(c, 'cc_analista2', rol='calidad', pin=PIN),
        'revisor': crear_usuario(c, 'cc_revisor', rol='calidad', pin=PIN),
        'jefe': crear_usuario(c, 'cc_jefe', rol='calidad', pin=PIN,
                              facultades='liberar_lotes,aprobar_maestros,gestionar_desviaciones'),
        'jefe2': crear_usuario(c, 'cc_jefe2', rol='calidad', pin=PIN,
                               facultades='liberar_lotes,aprobar_maestros,gestionar_desviaciones'),
        'almacen': crear_usuario(c, 'cc_almacen', rol='almacen', pin=PIN),
    }
    c.close()
    return g


@pytest.fixture
def conn():
    c = conexion()
    yield c
    c.rollback()
    c.close()


PRUEBAS = [
    {'prueba': 'Valoración', 'tipo': 'NUMERICA', 'limite_min': 98.0, 'limite_max': 102.0, 'unidad': '%',
     'decimales': 1, 'metodo': 'MA-012', 'en_certificado': 1},
    {'prueba': 'Aspecto', 'tipo': 'CUALITATIVA', 'texto_especificacion': 'Polvo fino blanco'},
    {'prueba': 'Interna', 'tipo': 'NUMERICA', 'limite_max': 5, 'decimales': 0, 'en_certificado': 0},
]


def vigente(conn, gente, material_id, pruebas=None, **datos):
    eid = esp.crear(conn, material_id=material_id, usuario='pruebas', datos=datos,
                    pruebas_lista=pruebas or PRUEBAS)
    esp.enviar_revision(conn, eid, usuario='pruebas')
    esp.revisar(conn, eid, firma=F(gente['revisor']))
    esp.aprobar(conn, eid, firma=F(gente['jefe']))
    conn.commit()
    return eid


def pruebas_de(conn, eid):
    return {p['prueba']: p['id'] for p in esp.pruebas(conn, eid)}


def analisis_completo(conn, gente, lote_id, eid, revisar=True, **muestreo):
    """Muestrea, captura todo conforme y (opcional) revisa. Devuelve el id del análisis."""
    aid = an.muestrear(conn, lote_id=lote_id, usuario='pruebas', firma=F(gente['analista']), **muestreo)
    ids = pruebas_de(conn, eid)
    for nombre, pid in ids.items():
        p = conn.execute('SELECT tipo FROM cc_especificacion_pruebas WHERE id = ?', (pid,)).fetchone()
        if p['tipo'] == 'NUMERICA':
            valor = '100.0' if nombre == 'Valoración' else '1'
            an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor=valor,
                        firma=F(gente['analista']))
        else:
            an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', texto_obs='Polvo fino blanco',
                        cumple='1', firma=F(gente['analista']))
    if revisar:
        an.revisar(conn, analisis_id=aid, firma=F(gente['revisor']))
    conn.commit()
    return aid


def mp_con_esp(conn, gente, **datos):
    mat = crear_material(conn, tipo='MP')
    return mat, vigente(conn, gente, mat, **datos)


def sin_bloqueos_de_prueba(conn):
    conn.execute('DELETE FROM prb_cosas WHERE bloquea = 1')
    conn.commit()


# 1 -----------------------------------------------------------------------------------------------

def test_especificacion_necesita_reviso_y_aprobo_de_personas_distintas(conn, gente):
    mat = crear_material(conn, tipo='MP')
    eid = esp.crear(conn, material_id=mat, usuario='pruebas', pruebas_lista=PRUEBAS)
    with pytest.raises(ErrorCalidad):
        esp.aprobar(conn, eid, firma=F(gente['jefe']))          # sin enviar ni revisar
    esp.enviar_revision(conn, eid, usuario='pruebas')
    with pytest.raises(ErrorCalidad, match='revisó'):
        esp.aprobar(conn, eid, firma=F(gente['jefe']))          # falta Revisó
    esp.revisar(conn, eid, firma=F(gente['jefe']))
    with pytest.raises(firmas.FirmaInvalida):
        esp.aprobar(conn, eid, firma=F(gente['jefe']))          # misma persona
    with pytest.raises(firmas.FirmaInvalida):
        esp.aprobar(conn, eid, firma=F(gente['revisor']))       # sin aprobar_maestros
    esp.aprobar(conn, eid, firma=F(gente['jefe2']))
    assert esp.especificacion(conn, eid)['estatus'] == 'VIGENTE'


def test_version_2_deja_obsoleta_la_1_y_el_analisis_abierto_conserva_su_version(conn, gente):
    mat, v1 = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat)
    aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista']))
    v2 = esp.nueva_version(conn, v1, usuario='pruebas')
    with pytest.raises(ErrorCalidad, match='qué cambió'):
        esp.enviar_revision(conn, v2, usuario='pruebas')
    esp.guardar(conn, v2, usuario='pruebas', datos={'cambios': 'Se amplió el límite'})
    esp.enviar_revision(conn, v2, usuario='pruebas')
    esp.revisar(conn, v2, firma=F(gente['revisor']))
    esp.aprobar(conn, v2, firma=F(gente['jefe']))
    assert esp.especificacion(conn, v1)['estatus'] == 'OBSOLETA'
    assert esp.vigente(conn, mat)['id'] == v2
    assert an.analisis(conn, aid)['especificacion_id'] == v1
    with pytest.raises(ErrorCalidad):
        esp.guardar(conn, v2, usuario='pruebas', datos={'referencia': 'otra'})   # la vigente no se edita


# 2 -----------------------------------------------------------------------------------------------

def test_bandeja_muestra_lotes_en_cuarentena_pero_no_presentaciones_de_un_granel(conn, gente):
    mat, _ = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat)
    gr = crear_material(conn, tipo='GR')
    pt = crear_material(conn, tipo='PT', unidad='pza', material_padre_id=gr)
    granel = crear_lote(conn, gr, origen='PRODUCCION')
    hijo = crear_lote(conn, pt, origen='PRODUCCION', lote_padre_id=granel)
    ids = {f['lote_id']: f for f in an.por_muestrear(conn)}
    assert lote in ids and granel in ids and hijo not in ids
    assert ids[lote]['especificacion_id'] and ids[granel]['especificacion_id'] is None   # "Falta especificación"
    with pytest.raises(ErrorCalidad, match='especificación vigente'):
        an.muestrear(conn, lote_id=granel, usuario='pruebas', firma=F(gente['analista']))
    # Sin especificación sí se puede rechazar (material que llegó dañado)
    dic.rechazar(conn, lote_id=granel, motivo='Llegó dañado', firma=F(gente['jefe']), usuario='pruebas')
    assert inventario.lote(conn, hijo)['estatus'] == 'RECHAZADO'


# 3 -----------------------------------------------------------------------------------------------

def test_muestreo_crea_el_analisis_y_descuenta_la_muestra(conn, gente):
    mat, _ = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat, cantidad=50)
    aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista']), cantidad=250,
                       en_gramos=True, contenedores='3 de 12')
    a = an.analisis(conn, aid)
    assert a['folio'].startswith(f'AN-{date.today().year}-') and a['cantidad_muestra'] == 0.25
    mov = conn.execute('SELECT * FROM lote_movimientos WHERE id = ?', (a['movimiento_id'],)).fetchone()
    assert mov['tipo'] == 'MUESTRA' and mov['folio'] == a['folio'] and mov['cantidad'] == -0.25
    assert inventario.lote(conn, lote)['existencia'] == pytest.approx(49.75)
    with pytest.raises(ErrorCalidad, match='ya tiene el análisis'):
        an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista2']), tipo='ESPECIAL')


def test_muestreo_en_cero_no_mueve_inventario_y_la_muestra_sale_de_la_presentacion(conn, gente):
    gr = crear_material(conn, tipo='GR')
    pt = crear_material(conn, tipo='PT', unidad='kg', material_padre_id=gr)
    vigente(conn, gente, gr)
    granel = crear_lote(conn, gr, origen='PRODUCCION', cantidad=10)
    hijo = crear_lote(conn, pt, origen='PRODUCCION', lote_padre_id=granel, cantidad=20)
    movs = conn.execute('SELECT COUNT(*) FROM lote_movimientos').fetchone()[0]
    a1 = an.muestrear(conn, lote_id=granel, usuario='pruebas', firma=F(gente['analista']), cantidad=0)
    assert conn.execute('SELECT COUNT(*) FROM lote_movimientos').fetchone()[0] == movs
    an.cancelar(conn, analisis_id=a1, motivo='Prueba', firma=F(gente['analista']))
    a2 = an.muestrear(conn, lote_id=granel, usuario='pruebas', firma=F(gente['analista']), cantidad=1,
                      muestra_lote_id=hijo)
    assert an.analisis(conn, a2)['lote_id'] == granel
    assert inventario.lote(conn, hijo)['existencia'] == 19 and inventario.lote(conn, granel)['existencia'] == 10


# 4 -----------------------------------------------------------------------------------------------

def test_redondeo_de_farmacopea_en_los_limites(conn, gente):
    assert str(redondear('97.95', 1)) == '98.0' and str(redondear('97.94', 1)) == '97.9'
    assert str(redondear('0.125', 2)) == '0.13'      # round() de Python daría 0.12
    mat, eid = mp_con_esp(conn, gente, pruebas=[PRUEBAS[0]])
    pid = pruebas_de(conn, eid)['Valoración']
    esperado = {'98.04': True, '97.95': True, '97.94': False}
    for valor, cumple in esperado.items():
        lote = crear_lote(conn, mat)
        aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista']))
        r = an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor=valor,
                        firma=F(gente['analista']))
        assert r['cumple'] is cumple, valor
        assert (r['investigacion_id'] is None) is cumple


def test_equipo_no_apto_bloquea_la_captura(conn, gente):
    prueba = dict(PRUEBAS[0], equipo_tipo='Balanza analítica')
    mat, eid = mp_con_esp(conn, gente, pruebas=[prueba])
    vencido = crear_equipo(conn, tipo='Balanza analítica', vence_calibracion='2020-01-01')
    bueno = crear_equipo(conn, tipo='Balanza analítica')
    otro = crear_equipo(conn, tipo='Potenciómetro')
    aid = an.muestrear(conn, lote_id=crear_lote(conn, mat), usuario='pruebas', firma=F(gente['analista']))
    pid = pruebas_de(conn, eid)['Valoración']
    captura = dict(analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='100', firma=F(gente['analista']))
    with pytest.raises(ErrorCalidad, match='pide equipo'):
        an.capturar(conn, **captura)
    with pytest.raises(ErrorCalidad, match='no está apto'):
        an.capturar(conn, equipo_id=vencido, **captura)
    with pytest.raises(ErrorCalidad):
        an.capturar(conn, equipo_id=otro, **captura)
    assert vencido not in [e['id'] for e in an.equipos_aptos(conn, 'Balanza analítica')]
    assert an.capturar(conn, equipo_id=bueno, **captura)['cumple']


def test_certificado_del_proveedor_solo_de_proveedor_aprobado_y_nunca_en_granel(conn, gente):
    mat, eid = mp_con_esp(conn, gente, pruebas=[PRUEBAS[0]])
    pid = pruebas_de(conn, eid)['Valoración']
    en_eval = inventario.crear_proveedor(conn, nombre='Proveedor en evaluación', usuario='pruebas')
    aprobado = inventario.crear_proveedor(conn, nombre='Proveedor aprobado', usuario='pruebas', estatus='Aprobado')
    for prov, debe_pasar in ((en_eval, False), (aprobado, True)):
        lote = crear_lote(conn, mat, proveedor_id=prov)
        aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista']))
        captura = dict(analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99', informe='CoA-1',
                       fuente='Certificado del proveedor', firma=F(gente['analista']))
        if debe_pasar:
            assert an.capturar(conn, **captura)['cumple']
        else:
            with pytest.raises(ErrorCalidad, match='no es un proveedor Aprobado'):
                an.capturar(conn, **captura)
    gr = crear_material(conn, tipo='GR')
    eg = vigente(conn, gente, gr, pruebas=[PRUEBAS[0]])
    aid = an.muestrear(conn, lote_id=crear_lote(conn, gr, origen='PRODUCCION'), usuario='pruebas',
                       firma=F(gente['analista']))
    with pytest.raises(ErrorCalidad, match='granel'):
        an.capturar(conn, analisis_id=aid, prueba_id=pruebas_de(conn, eg)['Valoración'], usuario='pruebas',
                    valor='99', fuente='Certificado del proveedor', firma=F(gente['analista']))


def test_externo_pide_laboratorio_informe_y_archivo(conn, gente):
    mat, eid = mp_con_esp(conn, gente, pruebas=[dict(PRUEBAS[0], fuente_default='Externo')])
    aid = an.muestrear(conn, lote_id=crear_lote(conn, mat), usuario='pruebas', firma=F(gente['analista']))
    pid = pruebas_de(conn, eid)['Valoración']
    with pytest.raises(ErrorCalidad, match='laboratorio'):
        an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99', firma=F(gente['analista']))
    lab = an.enviar_externo(conn, analisis_id=aid, usuario='pruebas', laboratorio_nuevo='Lab Externo SA')
    assert conn.execute('SELECT tipo FROM proveedores WHERE id = ?', (lab,)).fetchone()['tipo'] == 'Servicios'
    with pytest.raises(ErrorCalidad, match='informe'):
        an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99', firma=F(gente['analista']))
    with pytest.raises(ErrorCalidad, match='Adjunta'):
        an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99', informe='INF-9',
                    firma=F(gente['analista']))
    pdf = FileStorage(io.BytesIO(b'%PDF-1.4 informe'), filename='informe.pdf')
    r = an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99', informe='INF-9',
                    archivo=pdf, firma=F(gente['analista']))
    assert r['cumple'] and an.detalle(conn, aid)['archivos']


# 5 -----------------------------------------------------------------------------------------------

def _oos(conn, gente):
    mat, eid = mp_con_esp(conn, gente, pruebas=[PRUEBAS[0]])
    lote = crear_lote(conn, mat)
    aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['analista']))
    pid = pruebas_de(conn, eid)['Valoración']
    r = an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='90',
                    firma=F(gente['analista']))
    return lote, aid, pid, r


def test_resultado_fuera_de_especificacion_abre_investigacion_y_bloquea_la_revision(conn, gente):
    lote, aid, pid, r = _oos(conn, gente)
    assert not r['cumple'] and r['investigacion_id']
    assert inv.investigacion(conn, r['investigacion_id'])['folio'].startswith('OOS-')
    assert an.analisis(conn, aid)['estatus'] == 'EN_ANALISIS'
    with pytest.raises(ErrorCalidad, match='investigaciones abiertas'):
        an.revisar(conn, analisis_id=aid, firma=F(gente['revisor']))
    with pytest.raises(ErrorCalidad, match='no se corrige'):
        an.corregir(conn, resultado_id=r['resultado_id'], motivo='x', usuario='p', valor='99',
                    firma=F(gente['analista']))
    assert any('investigaciones' in m.lower() for m in dic.verificaciones(conn, lote)['motivos'])


def test_error_de_laboratorio_anula_el_resultado_y_pide_capturar_otra_vez(conn, gente):
    lote, aid, pid, r = _oos(conn, gente)
    iid = r['investigacion_id']
    revision = {c: '1' for c, _ in inv.REVISION}
    with pytest.raises(ErrorCalidad, match='marca'):
        inv.concluir(conn, iid, conclusion='ERROR_LABORATORIO', revision=revision, hallazgos='x',
                     firma=F(gente['analista']))
    revision['calculos_ok'] = '0'
    inv.concluir(conn, iid, conclusion='ERROR_LABORATORIO', revision=revision, hallazgos='Error de cálculo',
                 firma=F(gente['analista']))
    with pytest.raises(firmas.FirmaInvalida):
        inv.aprobar(conn, iid, firma=F(gente['revisor']))          # sin gestionar_desviaciones
    inv.aprobar(conn, iid, firma=F(gente['jefe']))
    res = conn.execute('SELECT * FROM cc_resultados WHERE id = ?', (r['resultado_id'],)).fetchone()
    assert res['valido'] == 0 and res['motivo_anulacion'].startswith('OOS-')
    assert [p['id'] for p in an.pendientes_de_capturar(conn, an.analisis(conn, aid))] == [pid]
    an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='100', firma=F(gente['analista']))
    assert an.analisis(conn, aid)['estatus'] == 'POR_REVISAR'
    assert len(an.detalle(conn, aid)['filas'][0]['resultados']) == 2     # la repetición queda a la vista


def test_oos_confirmado_deja_el_analisis_no_conforme(conn, gente):
    lote, aid, pid, r = _oos(conn, gente)
    revision = {c: '1' for c, _ in inv.REVISION}
    inv.concluir(conn, r['investigacion_id'], conclusion='CONFIRMADO', revision=revision,
                 hallazgos='Sin error de laboratorio', firma=F(gente['analista']))
    inv.aprobar(conn, r['investigacion_id'], firma=F(gente['jefe']))
    assert an.analisis(conn, aid)['estatus'] == 'NO_CONFORME'
    with pytest.raises(ErrorCalidad):
        dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    dic.rechazar(conn, lote_id=lote, motivo='OOS confirmado', firma=F(gente['jefe']), usuario='pruebas')
    assert inventario.lote(conn, lote)['estatus'] == 'RECHAZADO'


# 6 -----------------------------------------------------------------------------------------------

def test_corregir_captura_solo_antes_de_revisar_y_solo_quien_capturo(conn, gente):
    mat, eid = mp_con_esp(conn, gente, pruebas=[PRUEBAS[0]])
    aid = an.muestrear(conn, lote_id=crear_lote(conn, mat), usuario='pruebas', firma=F(gente['analista']))
    pid = pruebas_de(conn, eid)['Valoración']
    r = an.capturar(conn, analisis_id=aid, prueba_id=pid, usuario='pruebas', valor='99.1', firma=F(gente['analista']))
    with pytest.raises(ErrorCalidad, match='quien capturó'):
        an.corregir(conn, resultado_id=r['resultado_id'], motivo='dedo', usuario='p', valor='99.2',
                    firma=F(gente['analista2']))
    with pytest.raises(ErrorCalidad, match='motivo'):
        an.corregir(conn, resultado_id=r['resultado_id'], motivo='', usuario='p', valor='99.2',
                    firma=F(gente['analista']))
    nuevo = an.corregir(conn, resultado_id=r['resultado_id'], motivo='Error de dedo', usuario='p', valor='99.2',
                        firma=F(gente['analista']))
    viejo = conn.execute('SELECT * FROM cc_resultados WHERE id = ?', (r['resultado_id'],)).fetchone()
    assert viejo['valido'] == 0 and 'Error de captura' in viejo['motivo_anulacion']
    with pytest.raises(Exception):
        conn.execute('UPDATE cc_resultados SET valor = 1 WHERE id = ?', (nuevo['resultado_id'],))
    with pytest.raises(firmas.FirmaInvalida):
        an.revisar(conn, analisis_id=aid, firma=F(gente['analista']))        # quien capturó no revisa
    assert an.revisar(conn, analisis_id=aid, firma=F(gente['revisor'])) == 'CONFORME'
    with pytest.raises(ErrorCalidad, match='antes de la revisión'):
        an.corregir(conn, resultado_id=nuevo['resultado_id'], motivo='x', usuario='p', valor='99.3',
                    firma=F(gente['analista']))


# 7 -----------------------------------------------------------------------------------------------

def test_aprobar_falla_sin_analisis_conforme_caducado_o_sin_facultad(conn, gente):
    mat, eid = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat)
    with pytest.raises(ErrorCalidad, match='No hay análisis'):
        dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    aid = analisis_completo(conn, gente, lote, eid, revisar=False)
    with pytest.raises(ErrorCalidad, match='por revisar'):
        dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    an.revisar(conn, analisis_id=aid, firma=F(gente['revisor']))
    with pytest.raises(firmas.FirmaInvalida):
        dic.aprobar(conn, lote_id=lote, firma=F(gente['revisor']), usuario='pruebas')     # sin liberar_lotes
    ajena = dict(F(gente['jefe']), sesion={'username': 'otro', 'numero_empleado': gente['jefe2'], 'ip': None})
    with pytest.raises(firmas.FirmaInvalida, match='sesión'):
        dic.aprobar(conn, lote_id=lote, firma=ajena, usuario='pruebas')
    caducado = crear_lote(conn, mat, fecha_caducidad=(date.today() - timedelta(days=1)).isoformat())
    analisis_completo(conn, gente, caducado, eid)
    with pytest.raises(ErrorCalidad, match='Caducó'):
        dic.aprobar(conn, lote_id=caducado, firma=F(gente['jefe']), usuario='pruebas')
    dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    assert inventario.lote(conn, lote)['estatus'] == 'APROBADO'


def test_aprobar_falla_si_quien_dictamina_capturo_resultados(conn, gente):
    mat, eid = mp_con_esp(conn, gente, pruebas=[PRUEBAS[0]])
    lote = crear_lote(conn, mat)
    aid = an.muestrear(conn, lote_id=lote, usuario='pruebas', firma=F(gente['jefe']))
    an.capturar(conn, analisis_id=aid, prueba_id=pruebas_de(conn, eid)['Valoración'], usuario='pruebas',
                valor='100', firma=F(gente['jefe']))
    an.revisar(conn, analisis_id=aid, firma=F(gente['revisor']))
    with pytest.raises(ErrorCalidad, match='capturaste'):
        dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe2']), usuario='pruebas')


def test_aprobar_falla_con_la_verificacion_de_otro_modulo_y_muestra_el_motivo(conn, gente):
    mat, eid = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat)
    analisis_completo(conn, gente, lote, eid)
    conn.execute("INSERT INTO prb_cosas (nombre, bloquea) VALUES ('bloqueo', 1)")
    conn.commit()
    try:
        v = dic.verificaciones(conn, lote)
        assert 'Hay una cosa de prueba que bloquea' in v['modulos'] and not v['completa']
        with pytest.raises(ErrorCalidad, match='cosa de prueba'):
            dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
        # Rechazar y retener siempre se pueden, con motivo
        with pytest.raises(ErrorCalidad, match='motivo'):
            dic.retener(conn, lote_id=lote, motivo='', firma=F(gente['jefe']), usuario='pruebas')
        dic.retener(conn, lote_id=lote, motivo='En espera', firma=F(gente['jefe']), usuario='pruebas')
        assert inventario.lote(conn, lote)['estatus'] == 'RETENIDO'
        with pytest.raises(ErrorCalidad):
            dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')   # retenido: misma lista
    finally:
        conn.rollback()
        sin_bloqueos_de_prueba(conn)
    dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')


# 8 -----------------------------------------------------------------------------------------------

def test_aprobar_un_granel_aprueba_sus_presentaciones_y_guarda_la_lista(conn, gente):
    gr = crear_material(conn, tipo='GR')
    pt1 = crear_material(conn, tipo='PT', unidad='pza', material_padre_id=gr)
    pt2 = crear_material(conn, tipo='PT', unidad='pza', material_padre_id=gr)
    eid = vigente(conn, gente, gr, pide_retencion=1, cantidad_retencion=1)
    granel = crear_lote(conn, gr, origen='PRODUCCION', cantidad=100)
    h1 = crear_lote(conn, pt1, origen='PRODUCCION', lote_padre_id=granel, cantidad=10)
    h2 = crear_lote(conn, pt2, origen='PRODUCCION', lote_padre_id=granel, cantidad=10)
    analisis_completo(conn, gente, granel, eid)
    ret.guardar(conn, lote_id=h1, envases=2, cantidad=1, ubicacion='A1', usuario='p', firma=F(gente['analista']))
    with pytest.raises(ErrorCalidad, match='Falta la de'):
        dic.aprobar(conn, lote_id=granel, firma=F(gente['jefe']), usuario='pruebas')
    ret.guardar(conn, lote_id=h2, envases=2, cantidad=1, ubicacion='A1', usuario='p', firma=F(gente['analista']))
    afectados = dic.aprobar(conn, lote_id=granel, firma=F(gente['jefe']), usuario='pruebas')
    assert set(afectados) == {granel, h1, h2}
    assert all(inventario.lote(conn, x)['estatus'] == 'APROBADO' for x in (granel, h1, h2))
    d = dic.historial(conn, granel)[0]
    assert d['decision'] == 'APROBADO' and len(d['lotes_cascada']) == 2
    assert any(i['texto'].startswith('Muestras de retención') and i['ok'] for i in d['verificaciones']['calidad'])


# 9 -----------------------------------------------------------------------------------------------

def test_rechazar_un_lote_aprobado_lo_retiene_y_rechaza_con_una_sola_firma(conn, gente):
    mat, eid = mp_con_esp(conn, gente)
    lote = crear_lote(conn, mat)
    analisis_completo(conn, gente, lote, eid)
    dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    dic.rechazar(conn, lote_id=lote, motivo='Queja confirmada', firma=F(gente['jefe']), usuario='pruebas')
    hist = conn.execute("SELECT valor_anterior, valor_nuevo, firma_id FROM lote_historial WHERE lote_id = ? "
                        "AND campo = 'estatus' ORDER BY id DESC LIMIT 2", (lote,)).fetchall()
    assert [(h['valor_anterior'], h['valor_nuevo']) for h in reversed(hist)] == [('APROBADO', 'RETENIDO'),
                                                                                ('RETENIDO', 'RECHAZADO')]
    assert hist[0]['firma_id'] == hist[1]['firma_id']
    assert dic.historial(conn, lote)[0]['decision'] == 'RECHAZADO'


# 10 ----------------------------------------------------------------------------------------------

def test_reanalisis(conn, gente):
    mat, eid = mp_con_esp(conn, gente, meses_reanalisis=12)
    cad = (date.today() + timedelta(days=900)).isoformat()
    lote = crear_lote(conn, mat, fecha_caducidad=cad)
    analisis_completo(conn, gente, lote, eid)
    dic.aprobar(conn, lote_id=lote, firma=F(gente['jefe']), usuario='pruebas')
    assert inventario.lote(conn, lote)['fecha_reanalisis'] == sumar_meses(date.today().isoformat(), 12)
    assert lote not in [r['lote_id'] for r in an.reanalisis_proximos(conn)]
    pronto = (date.today() + timedelta(days=10)).isoformat()
    inventario.corregir_lote(conn, lote_id=lote, campo='fecha_reanalisis', valor=pronto, usuario='p',
                             motivo='prueba', firma_id=firmas.firmar(
                                 conn, numero_empleado=gente['jefe'], pin=PIN, significado='Liberó', modulo='calidad',
                                 entidad='lotes', entidad_id=lote, sesion=F(gente['jefe'])['sesion']))
    assert lote in [r['lote_id'] for r in an.reanalisis_proximos(conn)]
    with pytest.raises(ErrorCalidad, match='reanálisis CONFORME'):
        dic.extender_reanalisis(conn, lote_id=lote, nueva_fecha=cad, firma=F(gente['jefe']), usuario='p')
    analisis_completo(conn, gente, lote, eid, tipo='REANALISIS')
    despues = (date.today() + timedelta(days=901)).isoformat()
    with pytest.raises(ErrorCalidad, match='caducidad'):
        dic.extender_reanalisis(conn, lote_id=lote, nueva_fecha=despues, firma=F(gente['jefe']), usuario='p')
    dic.extender_reanalisis(conn, lote_id=lote, nueva_fecha=cad, firma=F(gente['jefe']), usuario='p')
    assert inventario.lote(conn, lote)['fecha_reanalisis'] == cad
    with pytest.raises(ErrorCalidad, match='reanálisis CONFORME'):          # cada extensión pide su reanálisis
        dic.extender_reanalisis(conn, lote_id=lote, nueva_fecha=pronto, firma=F(gente['jefe']), usuario='p')


# 11 ----------------------------------------------------------------------------------------------

def test_certificado_solo_de_aprobados_y_la_presentacion_usa_el_analisis_del_granel(conn, gente):
    gr = crear_material(conn, tipo='GR', registro_sanitario='Q-0000-001')
    pt = crear_material(conn, tipo='PT', unidad='pza', material_padre_id=gr)
    eid = vigente(conn, gente, gr)
    granel = crear_lote(conn, gr, origen='PRODUCCION', fecha_caducidad='2030-01-31')
    hijo = crear_lote(conn, pt, origen='PRODUCCION', lote_padre_id=granel, fecha_caducidad='2030-01-31')
    aid = an.muestrear(conn, lote_id=granel, usuario='pruebas', firma=F(gente['analista']))
    ids = pruebas_de(conn, eid)
    r = an.capturar(conn, analisis_id=aid, prueba_id=ids['Valoración'], usuario='p', valor='99.0',
                    firma=F(gente['analista']))
    an.corregir(conn, resultado_id=r['resultado_id'], motivo='dedo', usuario='p', valor='99.5',
                firma=F(gente['analista']))
    an.capturar(conn, analisis_id=aid, prueba_id=ids['Aspecto'], usuario='p', texto_obs='Polvo fino blanco',
                cumple='1', firma=F(gente['analista']))
    an.capturar(conn, analisis_id=aid, prueba_id=ids['Interna'], usuario='p', valor='2', firma=F(gente['analista']))
    an.revisar(conn, analisis_id=aid, firma=F(gente['revisor']))
    with pytest.raises(ErrorCalidad, match='aprobados'):
        cert.armar(conn, hijo)
    dic.aprobar(conn, lote_id=granel, firma=F(gente['jefe']), usuario='pruebas')
    c = cert.armar(conn, hijo)
    assert c['lote']['id'] == hijo and c['granel']['id'] == granel and c['analisis']['id'] == aid
    assert [p['prueba'] for p in c['pruebas']] == ['Valoración', 'Aspecto']          # Interna no va
    assert c['pruebas'][0]['resultado'] == '99.5 %'                                  # solo el válido
    assert c['reviso'] and c['libero'] and c['analizo']


# 12 ----------------------------------------------------------------------------------------------

def test_retencion_fecha_de_destruccion_y_envases(conn, gente):
    mat = crear_material(conn, tipo='MP')
    lote = crear_lote(conn, mat, cantidad=10, fecha_caducidad='2027-03-31')
    mid = ret.guardar(conn, lote_id=lote, envases=3, cantidad=0.5, ubicacion='B2', usuario='p',
                      firma=F(gente['analista']))
    m = ret.muestra(conn, mid)
    assert m['destruir_desde'] == '2028-03-31'
    assert inventario.lote(conn, lote)['existencia'] == 9.5
    with pytest.raises(ErrorCalidad, match='Solo quedan 3'):
        ret.usar(conn, muestra_id=mid, envases=4, motivo='Queja', firma=F(gente['analista']))
    ret.usar(conn, muestra_id=mid, envases=3, motivo='Queja', folio='DES-1', firma=F(gente['analista']))
    assert ret.muestra(conn, mid)['estatus'] == 'AGOTADA'
    with pytest.raises(ErrorCalidad, match='a partir del'):
        ret.destruir(conn, muestra_id=mid, firma=F(gente['analista']))
    sin_cad = crear_lote(conn, mat, cantidad=1)
    m2 = ret.muestra(conn, ret.guardar(conn, lote_id=sin_cad, envases=1, cantidad=0, ubicacion='B3', usuario='p',
                                       firma=F(gente['analista'])))
    assert m2['destruir_desde'] == sumar_meses(date.today().isoformat(), 12)


# 13 ----------------------------------------------------------------------------------------------

def test_analisis_por_equipo(conn, gente):
    prueba = dict(PRUEBAS[0], equipo_tipo='HPLC')
    mat, eid = mp_con_esp(conn, gente, pruebas=[prueba])
    eq = crear_equipo(conn, tipo='HPLC')
    aid = an.muestrear(conn, lote_id=crear_lote(conn, mat), usuario='p', firma=F(gente['analista']))
    an.capturar(conn, analisis_id=aid, prueba_id=pruebas_de(conn, eid)['Valoración'], usuario='p', valor='100',
                equipo_id=eq, firma=F(gente['analista']))
    hoy = date.today().isoformat()
    filas = registro.consulta('analisis_por_equipo', conn, eq, hoy)
    assert [f['analisis_id'] for f in filas] == [aid] and filas[0]['prueba'] == 'Valoración'
    # Equipos lee estos campos de cada renglón.
    assert {'lote_id', 'fecha', 'numero_lote', 'folio', 'material', 'url'} <= set(filas[0])
    assert filas[0]['fecha'] == hoy and filas[0]['url'] == f'/calidad/analisis/{aid}'
    manana = (date.today() + timedelta(days=1)).isoformat()
    assert registro.consulta('analisis_por_equipo', conn, eq, manana) == []


# 14 y pantallas ----------------------------------------------------------------------------------

def test_pantallas_y_dictamen_por_la_ruta(client, gente):
    c = conexion()
    sin_bloqueos_de_prueba(c)
    mat, eid = mp_con_esp(c, gente)
    lote = crear_lote(c, mat)
    aid = analisis_completo(c, gente, lote, eid)
    c.close()
    login(client, 'cc_jefe', PASSWORD)
    assert client.get('/calidad?tab=dictaminar').status_code == 200
    html = client.get(f'/calidad/lote/{lote}/dictamen').get_data(as_text=True)
    assert 'Listo para aprobar' in html or 'listo para aprobar' in html
    assert client.get(f'/calidad/analisis/{aid}').status_code == 200
    r = client.post(f'/calidad/lote/{lote}/aprobar', data={'firma_numero': gente['jefe'], 'firma_pin': PIN})
    assert r.status_code == 302
    c = conexion()
    assert inventario.lote(c, lote)['estatus'] == 'APROBADO'
    c.close()
    r = client.get(f'/calidad/lote/{lote}/certificado')
    assert r.status_code == 200 and 'Certificado de análisis' in r.get_data(as_text=True)
    assert client.get('/calidad/reportes/analisis.xlsx').data[:2] == b'PK'


def test_almacen_consulta_pero_no_opera(client, gente):
    login(client, 'cc_almacen', PASSWORD)
    assert client.get('/calidad').status_code == 200
    assert client.get('/calidad/especificaciones/carga').status_code == 302
    r = client.post('/calidad/lote/1/aprobar', data={'firma_numero': gente['almacen'], 'firma_pin': PIN})
    assert r.status_code == 302


def test_ganchos_registrados(gente):
    assert registro.modulo_instalado('calidad')
    assert registro.hay_consulta('analisis_por_equipo')
    with erp.app.test_request_context():
        assert registro.enlace('certificado', lote_id=5) == '/calidad/lote/5/certificado'
        assert registro.enlace('analisis', analisis_id=2) == '/calidad/analisis/2'
    c = conexion()
    usuario = {'numero_empleado': gente['jefe'], 'nombre': 'Jefe', 'username': 'cc_jefe',
               'roles': ['calidad'], 'facultades': ['liberar_lotes', 'aprobar_maestros']}
    with erp.app.app_context():
        textos = [p['texto'] for p in registro.recolectar('pendientes', c, usuario)]
        titulos = [k['titulo'] for k in registro.recolectar('kpis', c)]
    assert 'Lotes por dictaminar' in titulos and 'Investigaciones abiertas' in titulos
    assert isinstance(textos, list)
    c.close()


def test_especificaciones_desde_excel(conn, gente):
    from openpyxl import load_workbook
    mat = crear_material(conn, tipo='MP')
    codigo = conn.execute('SELECT codigo FROM materiales WHERE id = ?', (mat,)).fetchone()['codigo']
    wb = load_workbook(io.BytesIO(esp.plantilla_excel()))
    ws = wb.active
    for fila in ws.iter_rows(min_row=2):
        fila[0].value = codigo
    ws.append(['NO-EXISTE', '', '', '', '', '', '', '', 'pH', 'NUMERICA', 5, 7, '', 1])
    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    r = esp.importar_excel(conn, salida, usuario='pruebas')
    assert [c for c, _ in r['creadas']] == [codigo]
    assert any('NO-EXISTE' in e for e in r['errores'])
    eid = r['creadas'][0][1]
    e = esp.especificacion(conn, eid)
    assert e['estatus'] == 'BORRADOR' and e['meses_reanalisis'] == 24 and len(esp.pruebas(conn, eid)) == 2
    esp.enviar_revision(conn, eid, usuario='p')
    assert esp.en_bloque(conn, 'revisar', [eid], firma=F(gente['revisor'])) == [(eid, None)]
    assert esp.en_bloque(conn, 'aprobar', [eid], firma=F(gente['jefe'])) == [(eid, None)]
    assert esp.especificacion(conn, eid)['estatus'] == 'VIGENTE'
